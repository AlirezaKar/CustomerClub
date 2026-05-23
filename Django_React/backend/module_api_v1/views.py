import secrets as _secrets
import json
import time as _time
from pathlib import Path as _Path
from django.core.cache import cache
from rest_framework_simplejwt.authentication import JWTAuthentication
# from rest_framework.authentication import BasicAuthentication
from rest_framework.permissions import AllowAny, IsAuthenticated
from django.contrib.auth import get_user_model
from django.conf import settings
from django.utils import timezone
from django.db import IntegrityError, transaction
from rest_framework.response import Response
from rest_framework.request import Request
from rest_framework.views import APIView
from rest_framework import generics
from rest_framework import status
from rest_framework.exceptions import NotFound, PermissionDenied
from rest_framework_simplejwt.tokens import RefreshToken

from module_account import models as module_account_models
from module_account import serializers as module_account_serializers
from utils.functions import sms as sms_functions
from module_product import models as module_product_models
from module_product import serializers as module_product_serializers
from module_shop import models as module_shop_models
from module_shop import serializers as module_shop_serializers
from utils.card_uid import normalize_card_identifier
from utils.management_client import ManagementAPIError, management_configured
from utils.management_shops import fetch_management_shops
from utils.management_catalog import fetch_shop_catalog
from utils.management_card import link_end_user_card, unlink_end_user_card
from utils.management_redeem import redeem_special_offer
from utils.management_sync import ensure_local_shop, sync_customer_offers_for_user
from utils.management_transactions import record_special_offer_on_management
from utils.management_users import (
    lookup_end_user,
    lookup_end_user_by_card_uid,
    lookup_end_user_by_id,
    mirror_end_user_from_management,
    register_end_user,
)
from utils.stateless_auth import (
    KioskJWTAuthentication,
    issue_tokens_for_user_id,
)
from utils.stateless_kiosk import login_profile_for_user, stateless_kiosk_enabled
from utils.stateless_sms import send_login_code as stateless_send_login_code
from utils.stateless_sms import verify_login_code as stateless_verify_login_code

User= get_user_model()


def _parse_shop_id(raw) -> int | None:
    if raw in (None, ""):
        return None
    try:
        shop_id = int(raw)
    except (TypeError, ValueError):
        return None
    return shop_id if shop_id > 0 else None


_SIGNUP_CARD_ALREADY = "کارت قبلا ثبت شده است"
_SIGNUP_PHONE_ALREADY = "شماره تلفن قبلا ثبت شده است"


def _signup_error_message(exc: ManagementAPIError, *, has_card: bool) -> str:
    message = str(exc).strip()
    lower = message.lower()
    if "کارت" in message or "nfc" in lower or "uid" in lower:
        return _SIGNUP_CARD_ALREADY
    if "تلفن" in message or "phone" in lower or "username" in lower:
        return _SIGNUP_PHONE_ALREADY
    if exc.status_code == 409 and has_card:
        return _SIGNUP_CARD_ALREADY
    if exc.status_code == 409:
        return message or _SIGNUP_PHONE_ALREADY
    return message or "ثبت نام ناموفق بود."


def _local_signup_conflict_message(exc: IntegrityError, *, has_card: bool) -> str:
    text = str(exc).lower()
    if "card_uid" in text or (has_card and "unique" in text):
        return _SIGNUP_CARD_ALREADY
    if "phone_number" in text or "username" in text:
        return _SIGNUP_PHONE_ALREADY
    return _SIGNUP_PHONE_ALREADY if not has_card else _SIGNUP_CARD_ALREADY


def _management_shop_exists(shop_id: int) -> bool:
    shops = fetch_management_shops()
    if shops is None:
        return False
    return any(row.get("id") == shop_id for row in shops)


def _stateless_login_response(mgmt_user: dict, shop_id: int | None = None):
    user_data = login_profile_for_user(mgmt_user, shop_id=shop_id)
    tokens = issue_tokens_for_user_id(int(mgmt_user["id"]))
    return Response(
        {
            **tokens,
            "user": user_data,
        },
        status=status.HTTP_200_OK,
    )


def _login_response(user, shop_id: int | None = None):
    if stateless_kiosk_enabled():
        mgmt_user = {
            "id": user.id,
            "phone_number": getattr(user, "phone_number", "") or "",
            "score": getattr(user, "score", 0),
            "first_name": getattr(user, "first_name", "") or "",
            "last_name": getattr(user, "last_name", "") or "",
        }
        return _stateless_login_response(mgmt_user, shop_id=shop_id)

    sync_customer_offers_for_user(user, shop_id=shop_id)
    serialized_data = module_account_serializers.UserSerializer(
        user,
        many=False,
        context={"shop_id": shop_id},
    )
    refresh = RefreshToken.for_user(user)
    return Response(
        {
            "access": str(refresh.access_token),
            "refresh": str(refresh),
            "user": serialized_data.data,
        },
        status=status.HTTP_200_OK,
    )


def _agent_log(hypothesis_id, location, message, data):
    try:
        log_path = _Path(__file__).resolve().parents[2] / "debug-2b8592.log"
        with open(log_path, "a", encoding="utf-8") as f:
            f.write(json.dumps({
                "sessionId": "2b8592",
                "runId": "initial",
                "hypothesisId": hypothesis_id,
                "location": location,
                "message": message,
                "data": data,
                "timestamp": int(_time.time() * 1000),
            }) + "\n")
    except Exception:
        # ignore logging errors
        pass


def _masked_uid_from_raw(raw_uid: str) -> str:
    """Mask UID for logs (never log raw UID)."""
    try:
        hashed = module_account_models.NFCTag.hash_uid((raw_uid or "").strip().lower())
        return f"****{hashed[-4:]}"
    except Exception:
        return "****"


def _normalize_serializer_errors(errors):
    """
    Normalize errors to match this project's predominant API pattern:
    {"error": "...", ...}
    """
    if isinstance(errors, dict):
        if "error" in errors:
            return errors
        # DRF field errors: {"uid": ["This field is required."]}
        for field, msgs in errors.items():
            if isinstance(msgs, (list, tuple)) and msgs:
                return {"error": str(msgs[0]), "field": str(field)}
            return {"error": str(msgs), "field": str(field)}
    return {"error": "invalid request"}


def _require_jwt_user(request: Request):
    """
    Enforce JWT auth but return consistent JSON error shape.
    """
    try:
        auth = KioskJWTAuthentication() if stateless_kiosk_enabled() else JWTAuthentication()
        result = auth.authenticate(request)
        if not result:
            return None, Response({"error": "authentication required"}, status=status.HTTP_401_UNAUTHORIZED)
        user, _ = result
        request.user = user
        return user, None
    except Exception:
        return None, Response({"error": "authentication required"}, status=status.HTTP_401_UNAUTHORIZED)


class ShopLookupMixin:
    def _get_shop(self, request: Request, allow_public: bool= False):
        shop_id= request.GET.get("shop_id") or request.data.get("shop_id")
        user= getattr(request, "user", None)

        if user and user.is_authenticated:
            shop_qs= user.shop_set.all()
            if shop_id:
                shop= shop_qs.filter(id= shop_id).first()
            else:
                shop= shop_qs.first()
            if shop:
                return shop
            if not allow_public:
                raise PermissionDenied("Shop not found for current user.")

        if allow_public:
            public_qs= module_shop_models.Shop.objects.all()
            if shop_id:
                shop= public_qs.filter(id= shop_id).first()
            else:
                shop= public_qs.first()
            if shop:
                return shop

        raise NotFound("Shop not found.")


class CategoryLookupMixin:
    def _get_category(self, request: Request, category_id: int, allow_public: bool= False):
        base_qs= module_product_models.Category.objects.filter(id= category_id)
        user= getattr(request, "user", None)

        if user and user.is_authenticated:
            category= base_qs.filter(shop__owner= user).first()
            if category:
                return category
            if not allow_public:
                raise PermissionDenied("Category not found for current user.")

        if allow_public:
            category= base_qs.filter(is_active= True).first()
            if category:
                return category

        raise NotFound("Category not found.")


# class UserView(generics.ListCreateAPIView):
#     queryset= module_account_models.User.objects.all()
#     serializer_class= module_account_serializers.UserSerializer

# class UserDetailView(generics.RetrieveUpdateDestroyAPIView):
#     queryset= module_account_models.User.objects.all()
#     serializer_class= module_account_serializers.UserSerializer

# class UserView(APIView):
#     def get(self, request: Request):
#         # getting phone_number
#         phone_number= request.GET.get("phone_number")
#         if not phone_number:
#             return Response({"error": "required phone_number"}, status= status.HTTP_400_BAD_REQUEST)

#         # getting user
#         user= User.objects.filter(phone_number= phone_number).first()
#         if not user:
#             user= User.objects.create(username= phone_number, phone_number= phone_number)
#             user.set_unusable_password()
#             user.save()

#         # send login_verification_code if needed
#         send_login_verification_code= request.GET.get("send_login_verification_code")== "true"
#         if send_login_verification_code:
#             is_done, data_error= user.send_verification_code()
#             if is_done:
#                 return Response(None, status= status.HTTP_200_OK)
#             return Response({"time": data_error}, status= status.HTTP_425_TOO_EARLY)

#         # check login_verification_code
#         login_verification_code= request.GET.get("login_verification_code")
#         if login_verification_code:
#             user.update_verification_code()
#             if user.phone_number_verification_code== int(login_verification_code):
#                 user.update_verification_code(successful_use= True)
#                 serialized_data= module_account_serializers.UserSerializer(user, many= False)
#                 return Response(serialized_data.data, status= status.HTTP_200_OK)
#             return Response({"error": "login_verification_code is wrong"}, status= status.HTTP_400_BAD_REQUEST)

#         # submit user offer
#         user_offer_id= request.GET.get("user_offer_id")
#         if user_offer_id:
#             user_offer= module_product_models.Product.objects.filter(id= user_offer_id).first()
#             if user_offer:
#                 used_user_offer, created= user.useduseroffer_set.get_or_create(product= user_offer)
#                 if not created:
#                     used_user_offer.number+= 1
#                     used_user_offer.save()
#                 user.score-= user_offer.score
#                 user.save()
#                 return Response(None, status= status.HTTP_200_OK)

#                 # used_user_offer= user.useduseroffer_set.filter(product= user_offer)
#                 # if used_user_offer:
#                 #     used_user_offer.number+= 1
#                 #     used_user_offer.save()
#                 # else:
#                 #     used_user_offer= user.useduseroffer_set.create(product= user_offer)

#             return Response({"error": "user_offer_id is wrong"}, status= status.HTTP_400_BAD_REQUEST)


#         return Response({"error": "required send_login_verification_code or login_verification_code"}, status= status.HTTP_400_BAD_REQUEST)


class PublicShopsView(APIView):
    """List shops from the management server for kiosk shop selection."""

    authentication_classes = []
    permission_classes = [AllowAny]

    def get(self, request: Request):
        if not management_configured():
            return Response(
                {"error": "سرویس فروشگاه در حال حاضر در دسترس نیست."},
                status=status.HTTP_503_SERVICE_UNAVAILABLE,
            )

        shops = fetch_management_shops()
        if shops is None:
            return Response(
                {"error": "سرویس فروشگاه در حال حاضر در دسترس نیست."},
                status=status.HTTP_503_SERVICE_UNAVAILABLE,
            )

        return Response({"shops": shops}, status=status.HTTP_200_OK)


class UserSignup(APIView):
    def post(self, request: Request):
        phone_number = (request.data.get("phone_number") or "").strip()
        age = request.data.get("age")
        gender = request.data.get("gender")
        card_uid = normalize_card_identifier(
            (request.data.get("uid_code") or request.data.get("card_uid") or "").strip()
        ) or None

        if not phone_number:
            return Response({"error": "required phone_number"}, status=status.HTTP_400_BAD_REQUEST)
        if not age:
            return Response({"error": "required age"}, status=status.HTTP_400_BAD_REQUEST)
        if not gender:
            return Response({"error": "required gender"}, status=status.HTTP_400_BAD_REQUEST)

        if not management_configured():
            return Response(
                {"error": "سرویس ثبت نام در حال حاضر در دسترس نیست."},
                status=status.HTTP_503_SERVICE_UNAVAILABLE,
            )

        has_card = bool(card_uid)
        try:
            mgmt_user = register_end_user(
                phone_number=phone_number,
                age=int(age),
                gender=gender,
                card_uid=card_uid,
            )
        except ManagementAPIError as exc:
            if exc.status_code in (400, 409, 500):
                return Response(
                    {"error": _signup_error_message(exc, has_card=has_card)},
                    status=status.HTTP_400_BAD_REQUEST,
                )
            return Response(
                {"error": "سرویس ثبت نام در حال حاضر در دسترس نیست."},
                status=status.HTTP_503_SERVICE_UNAVAILABLE,
            )

        if not stateless_kiosk_enabled():
            try:
                mirror_end_user_from_management(mgmt_user)
            except IntegrityError as exc:
                return Response(
                    {"error": _local_signup_conflict_message(exc, has_card=has_card)},
                    status=status.HTTP_400_BAD_REQUEST,
                )
        return Response(None, status=status.HTTP_200_OK)


class UserLogin(APIView):
    def post(self, request: Request):
        card_uid = normalize_card_identifier(
            (request.data.get("uid_code") or request.data.get("card_uid") or "").strip()
        ) or None
        if card_uid:
            shop_id = _parse_shop_id(request.data.get("shop_id"))
            if shop_id is None:
                return Response({"error": "required shop_id"}, status=status.HTTP_400_BAD_REQUEST)

            if stateless_kiosk_enabled():
                if not _management_shop_exists(shop_id):
                    return Response({"error": "shop_not_found"}, status=status.HTTP_404_NOT_FOUND)
            elif ensure_local_shop(shop_id) is None:
                return Response({"error": "shop_not_found"}, status=status.HTTP_404_NOT_FOUND)

            if not management_configured():
                return Response(
                    {"error": "سرویس ورود در حال حاضر در دسترس نیست."},
                    status=status.HTTP_503_SERVICE_UNAVAILABLE,
                )

            try:
                mgmt_user = lookup_end_user_by_card_uid(card_uid)
            except ManagementAPIError:
                return Response(
                    {"error": "سرویس ورود در حال حاضر در دسترس نیست."},
                    status=status.HTTP_503_SERVICE_UNAVAILABLE,
                )

            if not mgmt_user:
                return Response(
                    {"error": "این کارت ثبت نشده است"},
                    status=status.HTTP_404_NOT_FOUND,
                )

            if stateless_kiosk_enabled():
                return _stateless_login_response(mgmt_user, shop_id=shop_id)

            user = mirror_end_user_from_management(mgmt_user)
            return _login_response(user, shop_id=shop_id)


        phone_number = (request.data.get("phone_number") or "").strip()
        shop_id = _parse_shop_id(request.data.get("shop_id"))
        if not phone_number:
            return Response({"error": "required phone_number"}, status=status.HTTP_400_BAD_REQUEST)
        if shop_id is None:
            return Response({"error": "required shop_id"}, status=status.HTTP_400_BAD_REQUEST)

        if stateless_kiosk_enabled():
            if not _management_shop_exists(shop_id):
                return Response({"error": "shop_not_found"}, status=status.HTTP_404_NOT_FOUND)
        elif ensure_local_shop(shop_id) is None:
            return Response({"error": "shop_not_found"}, status=status.HTTP_404_NOT_FOUND)

        if not management_configured():
            return Response(
                {"error": "سرویس ورود در حال حاضر در دسترس نیست."},
                status=status.HTTP_503_SERVICE_UNAVAILABLE,
            )

        try:
            mgmt_user = lookup_end_user(phone_number)
        except ManagementAPIError:
            return Response(
                {"error": "سرویس ورود در حال حاضر در دسترس نیست."},
                status=status.HTTP_503_SERVICE_UNAVAILABLE,
            )

        if not mgmt_user:
            return Response(
                {
                    "error": "کاربری با این شماره تلفن ثبت نشده است. لطفاً ابتدا ثبت نام کنید.",
                },
                status=status.HTTP_404_NOT_FOUND,
            )

        send_login_verification_code= request.data.get("send_login_verification_code")
        if send_login_verification_code:
            if stateless_kiosk_enabled():
                is_done, data_error = stateless_send_login_code(phone_number)
            else:
                user = mirror_end_user_from_management(mgmt_user)
                is_done, data_error= user.send_verification_code()
            if is_done:
                return Response(None, status= status.HTTP_200_OK)
            if data_error == "resend_limit_reached":
                return Response({"error": "شما به محدودیت تعداد ارسال پیامک برای این شماره تلفن رسیده است."}, status=status.HTTP_429_TOO_MANY_REQUESTS)
            return Response({"time": data_error}, status= status.HTTP_425_TOO_EARLY)

        login_verification_code= request.data.get("login_verification_code")
        if login_verification_code:
            if stateless_kiosk_enabled():
                if not stateless_verify_login_code(phone_number, str(login_verification_code)):
                    return Response({"error": "کد وارد شده نادرست است."}, status= status.HTTP_400_BAD_REQUEST)
                return _stateless_login_response(mgmt_user, shop_id=shop_id)

            user = mirror_end_user_from_management(mgmt_user)
            user.update_verification_code()
            if user.phone_number_verification_code== int(login_verification_code):
                user.update_verification_code(successful_use= True)
                try:
                    cache.delete(f"sms_resend:{user.phone_number}")
                except Exception:
                    pass
                return _login_response(user, shop_id=shop_id)
            return Response({"error": "کد وارد شده نادرست است."}, status= status.HTTP_400_BAD_REQUEST)

        return Response({"error": "required send_login_verification_code or login_verification_code"}, status= status.HTTP_400_BAD_REQUEST)


# NFC bridge login: hardware script POSTs UID here; we create a one-time token
# and return a URL for the browser to complete login (see NfcCompleteView).
class NfcLoginView(APIView):
    authentication_classes= []
    permission_classes= [AllowAny]

    def post(self, request: Request):
        uid = (request.data.get("uid") or request.data.get("card_uid") or "").strip()
        if not uid:
            return Response({"status": "failed", "message": "uid required"}, status= status.HTTP_400_BAD_REQUEST)

        secret = getattr(settings, "NFC_BRIDGE_SECRET", None)
        if secret:
            provided = request.data.get("secret") or request.headers.get("X-NFC-Bridge-Secret")
            if provided != secret:
                return Response({"status": "failed", "message": "unauthorized"}, status= status.HTTP_401_UNAUTHORIZED)

        user_id = None
        mgmt_user = None
        user = None
        if stateless_kiosk_enabled():
            if not management_configured():
                return Response({"status": "failed", "message": "Service unavailable"}, status=status.HTTP_503_SERVICE_UNAVAILABLE)
            try:
                mgmt_user = lookup_end_user_by_card_uid(uid.lower())
            except ManagementAPIError:
                return Response({"status": "failed", "message": "Service unavailable"}, status=status.HTTP_503_SERVICE_UNAVAILABLE)
            if not mgmt_user:
                return Response({"status": "failed", "message": "Unknown card"}, status=status.HTTP_401_UNAUTHORIZED)
            user_id = mgmt_user["id"]
        else:
            user = User.objects.filter(card_uid=uid).first()
            if not user:
                return Response({"status": "failed", "message": "Unknown card"}, status=status.HTTP_401_UNAUTHORIZED)
            user_id = user.id

        one_time_token = _secrets.token_urlsafe(32)
        cache_key = f"nfc_login:{one_time_token}"
        cache.set(cache_key, user_id, timeout= getattr(settings, "NFC_LOGIN_TOKEN_TTL", 60))

        frontend_base = getattr(settings, "FRONTEND_BASE_URL", "http://127.0.0.1:5002")
        login_url = f"{frontend_base.strip('/')}/?nfc_token={one_time_token}"
        display_name = ""
        if stateless_kiosk_enabled():
            display_name = (mgmt_user.get("phone_number") or mgmt_user.get("username") or str(user_id))
        else:
            display_name = user.username
        return Response({
            "status": "success",
            "user": display_name,
            "token": one_time_token,
            "login_url": login_url,
        })


# Browser exchanges one-time token (from NFC bridge redirect) for user profile.
class NfcCompleteView(APIView):
    authentication_classes= []
    permission_classes= [AllowAny]

    def post(self, request: Request):
        token = (request.data.get("token") or "").strip()
        if not token:
            return Response({"error": "token required"}, status= status.HTTP_400_BAD_REQUEST)

        cache_key = f"nfc_login:{token}"
        user_id = cache.get(cache_key)
        if user_id is None:
            return Response({"error": "Invalid or expired token"}, status= status.HTTP_401_UNAUTHORIZED)
        cache.delete(cache_key)

        shop_id = _parse_shop_id(request.data.get("shop_id"))

        if stateless_kiosk_enabled():
            if shop_id is not None and not _management_shop_exists(shop_id):
                return Response({"error": "shop_not_found"}, status=status.HTTP_404_NOT_FOUND)
            try:
                mgmt_user = lookup_end_user_by_id(int(user_id))
            except ManagementAPIError:
                return Response(
                    {"error": "سرویس ورود در حال حاضر در دسترس نیست."},
                    status=status.HTTP_503_SERVICE_UNAVAILABLE,
                )
            if not mgmt_user:
                return Response({"error": "User not found"}, status=status.HTTP_404_NOT_FOUND)
            user_data = login_profile_for_user(mgmt_user, shop_id=shop_id)
            return Response(user_data, status=status.HTTP_200_OK)

        user = User.objects.filter(id=user_id).first()
        if not user:
            return Response({"error": "User not found"}, status=status.HTTP_404_NOT_FOUND)

        if shop_id is not None and ensure_local_shop(shop_id) is None:
            return Response({"error": "shop_not_found"}, status=status.HTTP_404_NOT_FOUND)

        sync_customer_offers_for_user(user, shop_id=shop_id)
        serialized_data = module_account_serializers.UserSerializer(
            user,
            many=False,
            context={"shop_id": shop_id},
        )
        return Response(serialized_data.data, status=status.HTTP_200_OK)


class NfcRegisterView(APIView):
    """
    Link/register an NFC tag for the authenticated user.

    Requires JWT authentication (same pattern as other protected endpoints).
    """

    # We enforce JWT manually to keep error format consistent.
    authentication_classes = []
    permission_classes = [AllowAny]

    def post(self, request: Request):
        user, auth_resp = _require_jwt_user(request)
        if auth_resp is not None:
            _agent_log(
                "NFC_REGISTER_UNAUTH",
                "backend/module_api_v1/views.py:NfcRegisterView",
                "nfc_register_unauthorized",
                {"userId": None},
            )
            return auth_resp

        serializer = module_account_serializers.NfcRegisterSerializer(
            data=request.data,
            context={"request": request},
        )

        if not serializer.is_valid():
            payload = _normalize_serializer_errors(serializer.errors)
            _agent_log(
                "NFC_REGISTER_FAIL",
                "backend/module_api_v1/views.py:NfcRegisterView",
                "nfc_register_failed",
                {"userId": getattr(user, "id", None), "errors": payload},
            )
            return Response(payload, status=status.HTTP_400_BAD_REQUEST)

        raw_uid = (serializer.validated_data.get("uid") or "").strip().lower()

        if stateless_kiosk_enabled():
            try:
                link_end_user_card(user_id=user.id, card_uid=raw_uid)
            except ManagementAPIError as exc:
                status_code = status.HTTP_400_BAD_REQUEST
                if exc.status_code == 409:
                    status_code = status.HTTP_409_CONFLICT
                elif exc.status_code in (None, 503):
                    status_code = status.HTTP_503_SERVICE_UNAVAILABLE
                error = str(exc).strip()
                if exc.status_code == 409 or "کارت" in error or "nfc" in error.lower():
                    error = _SIGNUP_CARD_ALREADY
                return Response({"error": error}, status=status_code)
            masked = f"****{module_account_models.NFCTag.hash_uid(raw_uid)[-4:]}"
            return Response({"uid": masked, "tag_id": user.id}, status=status.HTTP_201_CREATED)

        hashed_uid = module_account_models.NFCTag.hash_uid(raw_uid)
        try:
            existing = module_account_models.NFCTag.objects.filter(uid=hashed_uid).first()
        except Exception as exc:
            _agent_log(
                "NFC_REGISTER_DB_ERR",
                "backend/module_api_v1/views.py:NfcRegisterView",
                "nfc_register_db_error",
                {"userId": user.id, "uid": f"****{hashed_uid[-4:]}", "error": str(exc)},
            )
            return Response({"error": "server_error"}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

        if existing and existing.user_id != user.id:
            _agent_log(
                "NFC_REGISTER_CONFLICT",
                "backend/module_api_v1/views.py:NfcRegisterView",
                "nfc_register_conflict_uid_linked_elsewhere",
                {"userId": user.id, "uid": f"****{hashed_uid[-4:]}"},
            )
            return Response(
                {"error": _SIGNUP_CARD_ALREADY, "field": "uid"},
                status=status.HTTP_409_CONFLICT,
            )

        try:
            tag = serializer.save()
        except Exception as exc:
            _agent_log(
                "NFC_REGISTER_SAVE_ERR",
                "backend/module_api_v1/views.py:NfcRegisterView",
                "nfc_register_save_error",
                {"userId": user.id, "uid": f"****{hashed_uid[-4:]}", "error": str(exc)},
            )
            return Response({"error": "server_error"}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

        _agent_log(
            "NFC_REGISTER_OK",
            "backend/module_api_v1/views.py:NfcRegisterView",
            "nfc_register_success",
            {"userId": user.id, "uid": tag.mask_uid(), "tagId": tag.id},
        )
        return Response({"uid": tag.mask_uid(), "tag_id": tag.id}, status=status.HTTP_201_CREATED)


class NfcAuthLoginView(APIView):
    """
    Standalone NFC login (no prior auth required).

    On success returns JWT tokens using SimpleJWT's existing token logic.
    """

    authentication_classes = []
    permission_classes = [AllowAny]

    def post(self, request: Request):
        # Rate limiting: per-IP counter to mitigate brute force (default 30/min).
        try:
            ip = request.META.get("HTTP_X_FORWARDED_FOR", "").split(",")[0].strip() or request.META.get("REMOTE_ADDR", "")
            window_seconds = int(getattr(settings, "NFC_LOGIN_RATE_WINDOW", 60))
            max_requests = int(getattr(settings, "NFC_LOGIN_RATE_MAX", 30))
            rate_key = f"nfc_rl:{ip}"
            current = cache.get(rate_key) or 0
            if current >= max_requests:
                retry_after = window_seconds
                _agent_log(
                    "NFC_LOGIN_RL",
                    "backend/module_api_v1/views.py:NfcAuthLoginView",
                    "nfc_login_rate_limited",
                    {"ip": ip, "retryAfter": retry_after},
                )
                resp = Response({"error": "rate_limited", "retry_after": retry_after}, status=status.HTTP_429_TOO_MANY_REQUESTS)
                resp["Retry-After"] = str(retry_after)
                return resp
            cache.set(rate_key, int(current) + 1, timeout=window_seconds)
        except Exception:
            # never block login on rate limiter failures
            pass

        shop_id = _parse_shop_id(request.data.get("shop_id"))
        if shop_id is None:
            return Response({"error": "required shop_id"}, status=status.HTTP_400_BAD_REQUEST)

        if stateless_kiosk_enabled():
            card_uid = normalize_card_identifier(
                (
                    request.data.get("uid_code")
                    or request.data.get("uid")
                    or request.data.get("card_uid")
                    or ""
                ).strip()
            ) or None
            if not card_uid:
                return Response(
                    _normalize_serializer_errors({"uid": ["uid required"]}),
                    status=status.HTTP_400_BAD_REQUEST,
                )
            if not _management_shop_exists(shop_id):
                return Response({"error": "shop_not_found"}, status=status.HTTP_404_NOT_FOUND)
            if not management_configured():
                return Response(
                    {"error": "سرویس ورود در حال حاضر در دسترس نیست."},
                    status=status.HTTP_503_SERVICE_UNAVAILABLE,
                )
            try:
                mgmt_user = lookup_end_user_by_card_uid(card_uid)
            except ManagementAPIError:
                return Response(
                    {"error": "سرویس ورود در حال حاضر در دسترس نیست."},
                    status=status.HTTP_503_SERVICE_UNAVAILABLE,
                )
            if not mgmt_user:
                return Response({"error": "invalid credentials"}, status=status.HTTP_401_UNAUTHORIZED)
            tokens_payload = {
                **issue_tokens_for_user_id(int(mgmt_user["id"])),
                "user": login_profile_for_user(mgmt_user, shop_id=shop_id),
            }
            _agent_log(
                "NFC_LOGIN_OK",
                "backend/module_api_v1/views.py:NfcAuthLoginView",
                "nfc_login_success",
                {"userId": mgmt_user["id"]},
            )
            return Response(tokens_payload, status=status.HTTP_200_OK)

        serializer = module_account_serializers.NfcLoginSerializer(
            data=request.data,
            context={"request": request},
        )

        if not serializer.is_valid():
            errors = _normalize_serializer_errors(serializer.errors)
            error_code = errors.get("error")

            if error_code == "nfc_tag_locked":
                raw_uid = (request.data.get("uid") or "").strip().lower()
                tag = None
                try:
                    hashed = module_account_models.NFCTag.hash_uid(raw_uid)
                    tag = module_account_models.NFCTag.objects.filter(uid=hashed).first()
                except Exception:
                    tag = None

                retry_after_seconds = None
                if tag and tag.locked_until:
                    retry_after_seconds = max(0, int((tag.locked_until - timezone.now()).total_seconds()))

                _agent_log(
                    "NFC_LOGIN_LOCKED",
                    "backend/module_api_v1/views.py:NfcAuthLoginView",
                    "nfc_login_locked",
                    {"uid": _masked_uid_from_raw(raw_uid), "retryAfter": retry_after_seconds},
                )

                payload = {"error": "nfc_tag_locked"}
                if retry_after_seconds is not None:
                    payload["retry_after"] = retry_after_seconds
                resp = Response(payload, status=status.HTTP_423_LOCKED)
                if retry_after_seconds is not None:
                    resp["Retry-After"] = str(retry_after_seconds)
                return resp

            if error_code == "invalid credentials":
                raw_uid = request.data.get("uid") or ""
                _agent_log(
                    "NFC_LOGIN_INVALID",
                    "backend/module_api_v1/views.py:NfcAuthLoginView",
                    "nfc_login_invalid_credentials",
                    {"uid": _masked_uid_from_raw(str(raw_uid))},
                )
                return Response({"error": "invalid credentials"}, status=status.HTTP_401_UNAUTHORIZED)

            _agent_log(
                "NFC_LOGIN_FAIL",
                "backend/module_api_v1/views.py:NfcAuthLoginView",
                "nfc_login_failed",
                {
                    "errors": errors,
                    "uid": _masked_uid_from_raw(str(request.data.get("uid") or "")),
                },
            )
            return Response(errors, status=status.HTTP_400_BAD_REQUEST)

        if ensure_local_shop(shop_id) is None:
            return Response({"error": "shop_not_found"}, status=status.HTTP_404_NOT_FOUND)
        try:
            user = serializer.login()
        except Exception as exc:
            raw_uid = request.data.get("uid") or ""
            _agent_log(
                "NFC_LOGIN_ERR",
                "backend/module_api_v1/views.py:NfcAuthLoginView",
                "nfc_login_server_error",
                {"uid": _masked_uid_from_raw(str(raw_uid)), "error": str(exc)},
            )
            return Response({"error": "server_error"}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)
        sync_customer_offers_for_user(user, shop_id=shop_id)
        refresh = RefreshToken.for_user(user)
        serialized_data = module_account_serializers.UserSerializer(
            user,
            many=False,
            context={"shop_id": shop_id},
        )
        tokens_payload = {
            "refresh": str(refresh),
            "access": str(refresh.access_token),
            "user": serialized_data.data,
        }

        _agent_log(
            "NFC_LOGIN_OK",
            "backend/module_api_v1/views.py:NfcAuthLoginView",
            "nfc_login_success",
            {"userId": user.id},
        )
        return Response(tokens_payload, status=status.HTTP_200_OK)


class NfcUnlinkView(APIView):
    """Deactivate (unlink) a tag owned by the authenticated user."""

    authentication_classes = []
    permission_classes = [AllowAny]

    def post(self, request: Request):
        user, auth_resp = _require_jwt_user(request)
        if auth_resp is not None:
            return auth_resp

        uid = (request.data.get("uid") or "").strip().lower()
        tag_id = request.data.get("tag_id")

        if not uid and not tag_id:
            return Response({"error": "uid or tag_id required"}, status=status.HTTP_400_BAD_REQUEST)

        if stateless_kiosk_enabled():
            try:
                unlink_end_user_card(user_id=user.id)
            except ManagementAPIError as exc:
                status_code = status.HTTP_503_SERVICE_UNAVAILABLE
                if exc.status_code == 404:
                    status_code = status.HTTP_404_NOT_FOUND
                return Response({"error": str(exc)}, status=status_code)
            return Response({"status": "success"}, status=status.HTTP_200_OK)

        tag = None
        try:
            if tag_id:
                tag = module_account_models.NFCTag.objects.filter(id=tag_id).first()
            else:
                hashed_uid = module_account_models.NFCTag.hash_uid(uid)
                tag = module_account_models.NFCTag.objects.filter(uid=hashed_uid).first()
        except Exception as exc:
            _agent_log(
                "NFC_UNLINK_DB_ERR",
                "backend/module_api_v1/views.py:NfcUnlinkView",
                "nfc_unlink_db_error",
                {"userId": user.id, "uid": _masked_uid_from_raw(uid), "error": str(exc)},
            )
            return Response({"error": "server_error"}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

        if not tag:
            return Response({"error": "nfc_tag_not_found"}, status=status.HTTP_400_BAD_REQUEST)

        if tag.user_id != user.id:
            return Response({"error": "forbidden"}, status=status.HTTP_403_FORBIDDEN)

        tag.is_active = False
        tag.save(update_fields=["is_active"])

        _agent_log(
            "NFC_UNLINK_OK",
            "backend/module_api_v1/views.py:NfcUnlinkView",
            "nfc_unlink_success",
            {"userId": user.id, "tagId": tag.id, "uid": tag.mask_uid()},
        )
        return Response({"status": "success"}, status=status.HTTP_200_OK)


class NfcStatusView(APIView):
    """Return NFC tags status for the authenticated user."""

    authentication_classes = []
    permission_classes = [AllowAny]

    def get(self, request: Request):
        user, auth_resp = _require_jwt_user(request)
        if auth_resp is not None:
            return auth_resp

        if stateless_kiosk_enabled():
            try:
                mgmt_user = lookup_end_user_by_id(user.id)
            except ManagementAPIError:
                return Response(
                    {"error": "server_error"},
                    status=status.HTTP_503_SERVICE_UNAVAILABLE,
                )
            card_uid = (mgmt_user or {}).get("card_uid")
            tags = []
            if card_uid:
                masked = f"****{module_account_models.NFCTag.hash_uid(card_uid.lower())[-4:]}"
                tags.append(
                    {
                        "id": user.id,
                        "uid": masked,
                        "is_active": True,
                        "created_at": "",
                        "last_used_at": None,
                        "failed_attempts": 0,
                        "locked_until": None,
                    }
                )
            return Response(
                {"has_active_nfc_tags": bool(tags), "tags": tags},
                status=status.HTTP_200_OK,
            )

        try:
            data = module_account_serializers.NfcStatusSerializer.from_user(user)
            return Response(data, status=status.HTTP_200_OK)
        except Exception as exc:
            _agent_log(
                "NFC_STATUS_ERR",
                "backend/module_api_v1/views.py:NfcStatusView",
                "nfc_status_server_error",
                {"userId": user.id, "error": str(exc)},
            )
            return Response({"error": "server_error"}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)


class UserSubmitOffer(APIView):
    def post(self, request: Request):
        phone_number = (request.data.get("phone_number") or "").strip()
        if not phone_number:
            return Response({"error": "required phone_number"}, status=status.HTTP_400_BAD_REQUEST)

        product_id = request.data.get("product_id")
        if not product_id:
            return Response({"error": "required product_id"}, status=status.HTTP_400_BAD_REQUEST)

        try:
            product_id = int(product_id)
        except (TypeError, ValueError):
            return Response({"error": "product_id is wrong"}, status=status.HTTP_400_BAD_REQUEST)

        if not management_configured():
            return Response(
                {"error": "سرویس ثبت پیشنهاد در حال حاضر در دسترس نیست."},
                status=status.HTTP_503_SERVICE_UNAVAILABLE,
            )

        if stateless_kiosk_enabled():
            try:
                mgmt_user = lookup_end_user(phone_number)
            except ManagementAPIError as exc:
                status_code = status.HTTP_503_SERVICE_UNAVAILABLE
                if exc.status_code in (None, 503):
                    status_code = status.HTTP_503_SERVICE_UNAVAILABLE
                return Response({"error": str(exc)}, status=status_code)
            if not mgmt_user:
                return Response(
                    {"error": "کاربری با این شماره تلفن ثبت نشده است."},
                    status=status.HTTP_404_NOT_FOUND,
                )
            try:
                payload = redeem_special_offer(
                    user_id=int(mgmt_user["id"]),
                    product_id=product_id,
                )
            except ManagementAPIError as exc:
                status_code = status.HTTP_400_BAD_REQUEST
                if exc.status_code == 404:
                    status_code = status.HTTP_404_NOT_FOUND
                elif exc.status_code == 403:
                    status_code = status.HTTP_403_FORBIDDEN
                elif exc.status_code in (None, 503):
                    status_code = status.HTTP_503_SERVICE_UNAVAILABLE
                return Response({"error": str(exc)}, status=status_code)

            updated_user = payload.get("user") or mgmt_user
            shop_id = _parse_shop_id(request.data.get("shop_id"))
            user_data = login_profile_for_user(updated_user, shop_id=shop_id)
            return Response(user_data, status=status.HTTP_200_OK)

        user = User.objects.filter(phone_number=phone_number).first()
        if not user:
            return Response(
                {"error": "کاربری با این شماره تلفن ثبت نشده است."},
                status=status.HTTP_404_NOT_FOUND,
            )

        product = module_product_models.Product.objects.filter(
            pk=product_id,
            is_active=True,
            is_offerable=True,
        ).first()
        if not product:
            return Response(
                {"error": "محصول یافت نشد یا قابل استفاده در پیشنهاد ویژه نیست."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        if not module_product_models.UserSpecialOffer.objects.filter(
            user=user,
            product=product,
        ).exists():
            return Response(
                {"error": "این پیشنهاد به شما تعلق ندارد."},
                status=status.HTTP_403_FORBIDDEN,
            )

        shop_id = product.shop_id
        offer_tx = None

        try:
            with transaction.atomic():
                locked_user = User.objects.select_for_update().get(pk=user.pk)
                if locked_user.score < product.score:
                    return Response(
                        {"error": "امتیاز کافی نیست."},
                        status=status.HTTP_400_BAD_REQUEST,
                    )

                locked_user.score -= int(product.score)
                locked_user.save(update_fields=["score"])

                offer_tx = module_product_models.SpecialOfferTransaction.objects.create(
                    user=locked_user,
                    delta=-int(product.score),
                    reason=module_product_models.SpecialOfferTransaction.REASON_SPECIAL_OFFER,
                    product=product,
                )
                user = locked_user
                user_score_after = locked_user.score

            try:
                record_special_offer_on_management(
                    user_id=user.pk,
                    product_id=product.pk,
                    user_score=user_score_after,
                )
            except ManagementAPIError as exc:
                with transaction.atomic():
                    if offer_tx is not None:
                        offer_tx.delete()
                    locked_user = User.objects.select_for_update().get(pk=user.pk)
                    locked_user.score += int(product.score)
                    locked_user.save(update_fields=["score"])
                    user = locked_user

                status_code = status.HTTP_400_BAD_REQUEST
                if exc.status_code == 404:
                    status_code = status.HTTP_404_NOT_FOUND
                elif exc.status_code in (None, 503):
                    status_code = status.HTTP_503_SERVICE_UNAVAILABLE
                return Response({"error": str(exc)}, status=status_code)
        except User.DoesNotExist:
            return Response(
                {"error": "کاربری با این شماره تلفن ثبت نشده است."},
                status=status.HTTP_404_NOT_FOUND,
            )

        serialized_data = module_account_serializers.UserSerializer(
            user,
            many=False,
            context={"shop_id": shop_id},
        )
        return Response(serialized_data.data, status=status.HTTP_200_OK)


class Shops(APIView):
    authentication_classes= [JWTAuthentication]
    permission_classes= [IsAuthenticated]

    def get(self, request: Request):
        # getting user
        user= request.user
        shops: list[module_shop_models.Shop]= user.shop_set.all()
        if not shops:
            return Response(None, status= status.HTTP_403_FORBIDDEN)

        serialized_data= module_shop_serializers.ShopSerializer(shops, many= True)
        return Response(serialized_data.data, status= status.HTTP_200_OK)


class ShopSettings(ShopLookupMixin, APIView):
    authentication_classes= [JWTAuthentication]
    permission_classes= [IsAuthenticated]

    def get(self, request: Request):
        shop= self._get_shop(request)
        serialized_data= module_shop_serializers.ShopSettingsSerializer(shop, many= False)
        return Response(serialized_data.data, status= status.HTTP_200_OK)

    def patch(self, request: Request):
        shop= self._get_shop(request)
        serializer= module_shop_serializers.ShopSettingsSerializer(shop, data= request.data, partial= True)
        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data, status= status.HTTP_200_OK)
        return Response(serializer.errors, status= status.HTTP_400_BAD_REQUEST)


class ShopSettingsScore(ShopLookupMixin, APIView):
    authentication_classes= [JWTAuthentication]
    permission_classes= [IsAuthenticated]

    def get(self, request: Request):
        # # getting user
        # user= request.user

        # # check if user is shop owner
        # shop: module_shop_models.Shop= user.shop_set.first()
        # if not shop:
        #     return Response(None, status= status.HTTP_403_FORBIDDEN)

        try:
            shop= self._get_shop(request)
        except NotFound:
            return Response({"error": "shop_not_found"}, status= status.HTTP_404_NOT_FOUND)


        score= getattr(shop, "score", None)
        if not score:
            score= module_shop_models.Score.objects.create(shop= shop)

        serialized_data= module_shop_serializers.ScoreSerializer(score, many= False)
        return Response(serialized_data.data, status= status.HTTP_200_OK)


class ShopSettingsProducts(ShopLookupMixin, APIView):
    authentication_classes= [JWTAuthentication]
    permission_classes= [AllowAny]

    def get(self, request: Request):
        try:
            shop= self._get_shop(request, allow_public= True)
        except NotFound:
            return Response({"error": "shop_not_found"}, status= status.HTTP_404_NOT_FOUND)

        products= shop.product_set.all()
        is_shop_owner= (
            request.user.is_authenticated
            and shop.owner_id == request.user.id
        )
        if not is_shop_owner:
            products= products.filter(is_active=True)

        serialized_data= module_product_serializers.ProductSerializer(products, many= True)
        return Response(serialized_data.data, status= status.HTTP_200_OK)

    def post(self, request: Request):
        if not request.user or not request.user.is_authenticated:
            return Response({"detail": "Authentication credentials were not provided."}, status= status.HTTP_401_UNAUTHORIZED)

        shop= self._get_shop(request)
        serializer= module_product_serializers.ProductAdminSerializer(data= request.data, context= {"shop": shop})
        if serializer.is_valid():
            product= serializer.save()
            read_serializer= module_product_serializers.ProductSerializer(product, many= False)
            return Response(read_serializer.data, status= status.HTTP_201_CREATED)
        return Response(serializer.errors, status= status.HTTP_400_BAD_REQUEST)


class ShopSettingsProductDetail(APIView):
    authentication_classes= [JWTAuthentication]
    permission_classes= [IsAuthenticated]

    def get_object(self, request: Request, product_id: int):
        product= module_product_models.Product.objects.filter(id= product_id, shop__owner= request.user).first()
        if not product:
            raise NotFound("Product not found.")
        return product

    def get(self, request: Request, product_id: int):
        product= self.get_object(request, product_id)
        serializer= module_product_serializers.ProductSerializer(product, many= False)
        return Response(serializer.data, status= status.HTTP_200_OK)

    def patch(self, request: Request, product_id: int):
        product= self.get_object(request, product_id)
        serializer= module_product_serializers.ProductAdminSerializer(product, data= request.data, partial= True, context= {"shop": product.shop})
        if serializer.is_valid():
            product= serializer.save()
            read_serializer= module_product_serializers.ProductSerializer(product, many= False)
            return Response(read_serializer.data, status= status.HTTP_200_OK)
        return Response(serializer.errors, status= status.HTTP_400_BAD_REQUEST)

    def delete(self, request: Request, product_id: int):
        product= self.get_object(request, product_id)
        product.delete()
        return Response(status= status.HTTP_204_NO_CONTENT)


class ShopCategories(ShopLookupMixin, APIView):
    authentication_classes= [JWTAuthentication]
    permission_classes= [AllowAny]

    def get(self, request: Request):
        shop_id = _parse_shop_id(request.GET.get("shop_id"))
        if stateless_kiosk_enabled():
            if shop_id is None:
                return Response({"error": "required shop_id"}, status=status.HTTP_400_BAD_REQUEST)
            if not _management_shop_exists(shop_id):
                return Response({"error": "shop_not_found"}, status=status.HTTP_404_NOT_FOUND)
            catalog = fetch_shop_catalog(shop_id)
            if catalog is None:
                return Response(
                    {"error": "سرویس فروشگاه در حال حاضر در دسترس نیست."},
                    status=status.HTTP_503_SERVICE_UNAVAILABLE,
                )
            return Response(catalog, status=status.HTTP_200_OK)

        try:
            shop= self._get_shop(request, allow_public= True)
        except NotFound:
            return Response({"error": "shop_not_found"}, status= status.HTTP_404_NOT_FOUND)

        categories= shop.categories.all()
        is_shop_owner= (
            request.user.is_authenticated
            and shop.owner_id == request.user.id
        )
        if not is_shop_owner:
            categories= categories.filter(is_active=True)

        serializer= module_product_serializers.CategorySerializer(
            categories,
            many=True,
            context={"request": request, "shop": shop},
        )
        return Response(serializer.data, status= status.HTTP_200_OK)

    def post(self, request: Request):
        if not request.user or not request.user.is_authenticated:
            return Response({"detail": "Authentication credentials were not provided."}, status= status.HTTP_401_UNAUTHORIZED)

        shop= self._get_shop(request)
        serializer= module_product_serializers.CategoryDetailSerializer(data= request.data, context= {"shop": shop})
        if serializer.is_valid():
            category= serializer.save()
            read_serializer= module_product_serializers.CategoryDetailSerializer(category, context= {"shop": shop})
            return Response(read_serializer.data, status= status.HTTP_201_CREATED)
        return Response(serializer.errors, status= status.HTTP_400_BAD_REQUEST)


class ShopCategoryDetail(CategoryLookupMixin, APIView):
    authentication_classes= [JWTAuthentication]
    permission_classes= [AllowAny]

    def get(self, request: Request, category_id: int):
        category= self._get_category(request, category_id, allow_public= True)
        shop= category.shop
        serializer= module_product_serializers.CategorySerializer(
            category,
            many=False,
            context={"request": request, "shop": shop},
        )
        return Response(serializer.data, status= status.HTTP_200_OK)

    def patch(self, request: Request, category_id: int):
        if not request.user or not request.user.is_authenticated:
            return Response({"detail": "Authentication credentials were not provided."}, status= status.HTTP_401_UNAUTHORIZED)

        category= self._get_category(request, category_id)
        serializer= module_product_serializers.CategoryDetailSerializer(category, data= request.data, partial= True, context= {"shop": category.shop})
        if serializer.is_valid():
            category= serializer.save()
            read_serializer= module_product_serializers.CategoryDetailSerializer(category, context= {"shop": category.shop})
            return Response(read_serializer.data, status= status.HTTP_200_OK)
        return Response(serializer.errors, status= status.HTTP_400_BAD_REQUEST)

    def delete(self, request: Request, category_id: int):
        if not request.user or not request.user.is_authenticated:
            return Response({"detail": "Authentication credentials were not provided."}, status= status.HTTP_401_UNAUTHORIZED)

        category= self._get_category(request, category_id)
        category.delete()
        return Response(status= status.HTTP_204_NO_CONTENT)


class UpdateUserProperties(APIView):
    def post(self, request: Request):
        if stateless_kiosk_enabled():
            return Response(
                {"error": "این endpoint در حالت بدون پایگاه داده غیرفعال است."},
                status=status.HTTP_503_SERVICE_UNAVAILABLE,
            )
        try:
            # user: user phine number
            # properties: list of product second id and user score
            for phone_number, properties in request.data.items():
                phone_number= "0"+ phone_number
                user= User.objects.filter(phone_number= phone_number).first()
                if not user:
                    user= User.objects.create(username= phone_number, phone_number= phone_number)
                    user.set_unusable_password()
                    user.save()
                for user_offer in module_product_models.UserSpecialOffer.objects.filter(user= user):
                    user_offer.delete()
                for offer in properties["offers"]:
                    product= module_product_models.Product.objects.filter(
                        second_id=offer,
                        is_active=True,
                        is_offerable=True,
                    ).first()
                    if product:
                        offer_rate= random.randrange(100, 1000)
                        module_product_models.UserSpecialOffer.objects.create(user= user, product= product, offer_rate= offer_rate)
                earned_score = int(properties["score"])
                if earned_score:
                    user.score+= earned_score
                    user.save(update_fields=["score"])
                    module_product_models.SpecialOfferTransaction.objects.create(
                        user=user,
                        delta=earned_score,
                        reason=module_product_models.SpecialOfferTransaction.REASON_PURCHASE,
                    )

            return Response(None, status= status.HTTP_200_OK)
        except:
            return Response(None, status= status.HTTP_400_BAD_REQUEST)




# read data from excel (for test)
import openpyxl
import random
from os import listdir
from os.path import isfile, join
class AddItemsTest(APIView):
    def post(self, request: Request):
        if stateless_kiosk_enabled():
            return Response(
                {"error": "این endpoint در حالت بدون پایگاه داده غیرفعال است."},
                status=status.HTTP_503_SERVICE_UNAVAILABLE,
            )

        ##############################
        # add products
        item_dic_excel= openpyxl.load_workbook("item_dic.xlsx")
        item_dic_excel = item_dic_excel.active
        image_dir= r"E:\Python\Django\Aras\CustomerClub\media\temp\temp"
        # image_files = [(image_dir+ "\\"+ f).replace("\\", "/") for f in listdir(image_dir) if isfile(join(image_dir, f))]
        image_files = ["temp/temp/"+ f for f in listdir(image_dir) if isfile(join(image_dir, f))]
        # print(image_files)
        shop= module_shop_models.Shop.objects.all().first()

        for row in range(1, item_dic_excel.max_row):

            # col= 1: id, col= 2: name
            price_score= random.randrange(100, 1000)
            product= {
                "second_id": None,
                "title": None,
                "price": price_score,
                "score": price_score,
                "image": random.choice(image_files)
            }
            for col in item_dic_excel.iter_cols(1, item_dic_excel.max_column):
                if not product["second_id"]:
                    product["second_id"]= col[row].value
                elif not product["title"]:
                    product["title"]= col[row].value

            module_product_models.Product.objects.create(shop= shop, second_id= product["second_id"], title= product["title"], price= product["price"], score= product["score"], image= product["image"])
        ##############################

        ##############################
        # add user
        recom_Aras_transactions_excel= openpyxl.load_workbook("recom_Aras_transactions.xlsx")
        recom_Aras_transactions_excel = recom_Aras_transactions_excel.active

        # for row in range(1, recom_Aras_transactions_excel.max_row):
        for row in range(1, 51):
            # col= 1: phone_number, col= 2-5: product second_id
            user= {
                "phone_number": None,
                "products": [],
            }
            for col in recom_Aras_transactions_excel.iter_cols(1, recom_Aras_transactions_excel.max_column):
                if not user["phone_number"]:
                    user["phone_number"]= "0"+ str(col[row].value)
                else:
                    product= module_product_models.Product.objects.filter(second_id= col[row].value).first()
                    if product:
                        user["products"].append(product)

            new_user= User.objects.create(username= user["phone_number"], phone_number= user["phone_number"])
            new_user.set_unusable_password()
            new_user.save()

            for product in user["products"]:
                offer_rate= random.randrange(100, 1000)
                if product.is_offerable:
                    module_product_models.UserSpecialOffer.objects.create(user= new_user, product= product, offer_rate= offer_rate)
        ##############################

        return Response(None, status= status.HTTP_200_OK)
